import io
import re
from typing import List, Dict, Any, Optional
import pandas as pd
import openpyxl

from app.utils.text_cleaner import mask_placeholder, clean_vendor_name
from app.models.raw_data import RawIngestItem, IngestionResult

EXPECTED_HEADER_TOKENS = {
    "part", "mfg", "sku", "number", "num", "desc", "description",
    "brand", "vendor", "manufacturer", "manuf", "item", "upc", 
    "cat", "dept", "class", "fine", "unilog", "e1", "dib", "name"
}

def unmerge_and_flatten_excel(file_bytes: bytes) -> List[List[Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    # Unmerge ranges and replicate values across all bound cells
    for merge_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = merge_range.bounds
        top_left_val = ws.cell(row=min_row, column=min_col).value
        ws.unmerge_cells(str(merge_range))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ws.cell(row=row, column=col).value = top_left_val

    grid: List[List[Any]] = []
    for row in ws.iter_rows(values_only=True):
        grid.append(list(row))
    return grid

def score_header_row(row_cells: List[Any]) -> float:
    if not row_cells:
        return 0.0

    non_empty = [str(c).strip() for c in row_cells if c is not None and str(c).strip()]
    if not non_empty:
        return 0.0

    text_cells = [c for c in non_empty if not c.replace(".", "", 1).isdigit()]
    text_ratio = len(text_cells) / len(non_empty)

    token_matches = 0
    for cell in non_empty:
        words = re.findall(r"[a-zA-Z0-9]+", cell.lower())
        if any(w in EXPECTED_HEADER_TOKENS for w in words):
            token_matches += 1

    keyword_ratio = token_matches / len(non_empty)
    return (text_ratio * 0.3) + (keyword_ratio * 0.7)

def locate_header(grid: List[List[Any]], max_scan: int = 15) -> int:
    best_idx = 0
    highest_score = -1.0

    for i in range(min(len(grid), max_scan)):
        score = score_header_row(grid[i])
        if score > highest_score:
            highest_score = score
            best_idx = i

    return best_idx

def extract_primary_fields(cleaned_record: Dict[str, Optional[str]]) -> Dict[str, Optional[str]]:
    detected_mpn = None
    detected_mfr = None
    detected_desc = None

    # Priority-based token matching for core attributes
    for k, v in cleaned_record.items():
        k_clean = k.lower().replace(" ", "_").replace("-", "_")
        
        # MPN Detection
        if not detected_mpn and any(tag in k_clean for tag in ["mfg_part_num", "part_number", "part_num", "mpn", "sku", "item_num"]):
            detected_mpn = v

        # Manufacturer / Brand Detection
        if not detected_mfr and any(tag in k_clean for tag in ["part_manuf", "manufacturer", "manuf", "vendor", "brand", "e1_brand", "unilog_brand"]):
            detected_mfr = v

        # Description Detection
        if not detected_desc and any(tag in k_clean for tag in ["part_desc", "description", "desc", "item_desc", "product_name"]):
            detected_desc = v

    return {
        "mpn": detected_mpn,
        "mfr": detected_mfr,
        "desc": detected_desc
    }

def process_spatial_ingestion(file_bytes: bytes, filename: str) -> IngestionResult:
    if filename.endswith(".csv"):
        df_raw = pd.read_csv(io.BytesIO(file_bytes), header=None)
        grid = df_raw.values.tolist()
    else:
        grid = unmerge_and_flatten_excel(file_bytes)

    if not grid:
        return IngestionResult(
            filename=filename,
            total_rows=0,
            header_row_index=0,
            detected_columns=[],
            records=[]
        )

    header_idx = locate_header(grid)
    raw_headers = [str(c).strip() if c is not None and str(c).strip() != "" else f"COL_{idx}" for idx, c in enumerate(grid[header_idx])]

    # Deduplicate headers if identical column names exist
    seen_headers: Dict[str, int] = {}
    headers: List[str] = []
    for h in raw_headers:
        if h in seen_headers:
            seen_headers[h] += 1
            headers.append(f"{h}_{seen_headers[h]}")
        else:
            seen_headers[h] = 0
            headers.append(h)

    records: List[RawIngestItem] = []
    data_rows = grid[header_idx + 1:]

    for r_idx, row in enumerate(data_rows):
        # Ignore completely empty rows
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue

        raw_record: Dict[str, Any] = {}
        cleaned_record: Dict[str, Optional[str]] = {}

        for c_idx, col_name in enumerate(headers):
            val = row[c_idx] if c_idx < len(row) else None
            raw_record[col_name] = val
            cleaned_record[col_name] = mask_placeholder(val)

        fields = extract_primary_fields(cleaned_record)
        cleansed_vendor = clean_vendor_name(fields["mfr"])

        records.append(
            RawIngestItem(
                row_index=r_idx + 1,
                raw_record=raw_record,
                cleaned_record=cleaned_record,
                detected_mfr=cleansed_vendor or fields["mfr"],
                detected_mpn=fields["mpn"],
                detected_desc=fields["desc"],
                cleansed_vendor=cleansed_vendor
            )
        )

    return IngestionResult(
        filename=filename,
        total_rows=len(records),
        header_row_index=header_idx,
        detected_columns=headers,
        records=records
    )